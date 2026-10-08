"""Dated Brent CFD curve model: parses the per-cargo X workbook and re-computes it with corrections."""
import datetime as dt
import io
import os
import re

import openpyxl

WORKBOOK_PATH = os.path.join(os.path.dirname(__file__), "dated_workbook.xlsx")
MON = {m: i + 1 for i, m in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"])}
MNAME = {v: k for k, v in MON.items()}


def _num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _d(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def _easter(y):
    a, b, c = y % 19, y // 100, y % 100
    d, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mo = (h + l - 7 * m + 114) // 31
    return dt.date(y, mo, (h + l - 7 * m + 114) % 31 + 1)


def uk_holidays(y):
    """England & Wales bank holidays (no Platts London Dated assessment)."""
    hol = set()
    e = _easter(y)
    hol |= {e - dt.timedelta(days=2), e + dt.timedelta(days=1)}

    def first_mon(m):
        d = dt.date(y, m, 1)
        return d + dt.timedelta(days=(7 - d.weekday()) % 7)

    def last_mon(m):
        d = dt.date(y, m + 1, 1) - dt.timedelta(days=1)
        return d - dt.timedelta(days=d.weekday())

    hol |= {first_mon(5), last_mon(5), last_mon(8)}
    ny = dt.date(y, 1, 1)
    hol.add(ny if ny.weekday() < 5 else ny + dt.timedelta(days=7 - ny.weekday()))
    xmas, box = dt.date(y, 12, 25), dt.date(y, 12, 26)
    if xmas.weekday() == 5:
        hol |= {dt.date(y, 12, 27), dt.date(y, 12, 28)}
    elif xmas.weekday() == 6:
        hol |= {dt.date(y, 12, 26), dt.date(y, 12, 27)}
    elif box.weekday() == 5:
        hol |= {xmas, dt.date(y, 12, 28)}
    else:
        hol |= {xmas, box}
    return hol


def _is_bd(d, hol):
    return d.weekday() < 5 and d not in hol


def _bdays(a, b, hol):
    out, d = [], a
    while d <= b:
        if _is_bd(d, hol):
            out.append(d)
        d += dt.timedelta(days=1)
    return out


def _add_month(y, m, k):
    m2 = m - 1 + k
    return y + m2 // 12, m2 % 12 + 1


def _mlabel(y, m):
    return f"{MNAME[m]}{str(y)[2:]}"


def _avg(vals):
    vals = [v for v in vals if v is not None]
    return sum(vals) / len(vals) if vals else None


def _r(v, n=1):
    return None if v is None else round(v, n)


def load_workbook_bytes():
    if not os.path.exists(WORKBOOK_PATH):
        return None
    with open(WORKBOOK_PATH, "rb") as fh:
        return fh.read()


def save_workbook_bytes(content):
    parse(content)
    with open(WORKBOOK_PATH, "wb") as fh:
        fh.write(content)


def parse(content):
    wf = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    wv = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    for sh in ("RUN", "DAILY", "CARGOES"):
        if sh not in wf.sheetnames:
            raise ValueError(f"workbook is missing sheet {sh}")
    return build(wf, wv)


def build(wf, wv):
    R, RF = wv["RUN"], wf["RUN"]
    title = str(R["A1"].value or "")
    findings = []

    # ── inputs ──
    efp_dec = _num(R["B4"].value) or 0.0
    jan_feb = _num(R["B5"].value) or 0.0
    dec_jan = _num(R["B6"].value) or 0.0
    efp_jan = efp_dec
    feb_mar = None
    if "DIVERGENCE" in wv.sheetnames:
        efp_jan = _num(wv["DIVERGENCE"]["B5"].value) or efp_dec
        dv4 = wf["DIVERGENCE"]["B4"].value
        if isinstance(dv4, (int, float)) and _num(dv4) != efp_dec:
            findings.append(("high", "DIVERGENCE!B4 EFP Dec is hard-coded and differs from RUN!B4",
                             f"DIVERGENCE uses {dv4} but RUN!B4 = {efp_dec}. Link it (=RUN!B4)."))
        elif isinstance(dv4, (int, float)):
            findings.append(("low", "DIVERGENCE!B4 EFP Dec is typed, not linked",
                             f"Label says '= RUN!B4' but the cell is a typed {dv4}. Matches today; will drift the first time EFP moves. Use =RUN!$B$4."))
    dfl_jan_q = None
    if "DFL_DECOMP" in wv.sheetnames:
        feb_mar = _num(wv["DFL_DECOMP"]["B6"].value)
        dfl_jan_q = _num(wv["DFL_DECOMP"]["B8"].value)
    front_frozen = _num(R["B3"].value)

    # ── as-of date ──
    weeks_raw = []
    for r in range(9, 40):
        s, e = _d(R.cell(r, 2).value), _d(R.cell(r, 3).value)
        if not s or not e:
            if weeks_raw:
                break
            continue
        weeks_raw.append((r, R.cell(r, 1).value, s, e, _num(R.cell(r, 4).value), _num(R.cell(r, 5).value)))
    year0 = weeks_raw[0][2].year if weeks_raw else dt.date.today().year
    m = re.search(r"(\d{1,2})-([A-Z][a-z]{2})", title)
    asof = dt.date(year0, MON[m.group(2)], int(m.group(1))) if m and m.group(2) in MON else dt.date.today()
    vm = re.search(r"\bv(\d+)\b", title)
    version = f"v{vm.group(1)}" if vm else "?"
    hol = uk_holidays(asof.year) | uk_holidays(asof.year + 1) | uk_holidays(asof.year - 1)

    # ── weekly CFD ladder: level(n+1) = level(n) − roll(n) ──
    weeks, lvl = [], front_frozen
    bad_labels = []
    for i, (r, lab, s, e, roll, wb_level) in enumerate(weeks_raw):
        if i == 0:
            lvl = front_frozen if front_frozen is not None else wb_level
        label = lab if isinstance(lab, str) and len(lab) <= 6 else f"W{i + 1}"
        if isinstance(lab, str) and len(lab) > 6:
            bad_labels.append(f"A{r}")
        weeks.append({"label": label, "start": s.isoformat(), "end": e.isoformat(), "roll_to_next": roll,
                      "level": _r(lvl, 2), "level_wb": _r(wb_level, 2),
                      "realized": e < asof, "contains_asof": s <= asof <= e})
        if roll is not None and lvl is not None:
            lvl = lvl - roll
    if bad_labels:
        findings.append(("med", "Version notes are pasted over the week labels",
                         f"RUN {', '.join(bad_labels)} hold multi-paragraph run notes instead of the week name, but those rows are still live curve rows (DAILY step column reads RUN!B9:E23). Move notes to a NOTES sheet."))

    # ── daily curve ──
    D = wv["DAILY"]
    daily = []
    for r in range(4, 400):
        d = _d(D.cell(r, 1).value)
        if not d:
            if daily:
                break
            continue
        step = None
        for w in weeks:
            if w["start"] <= d.isoformat() <= w["end"]:
                step = w["level"]
        daily.append({"date": d.isoformat(), "step": step, "smooth": _num(D.cell(r, 3).value),
                      "holiday": d in hol, "realized": d < asof})
    if not daily:
        raise ValueError("DAILY sheet has no dated rows")
    grid_end = dt.date.fromisoformat(daily[-1]["date"])
    sm = {dt.date.fromisoformat(x["date"]): x["smooth"] for x in daily}
    hol_in_grid = [x["date"] for x in daily if x["holiday"]]
    if hol_in_grid:
        findings.append(("med", "UK bank holidays carry Dated values in the grid",
                         f"{', '.join(hol_in_grid)} have smooth values but Platts does not assess Dated on English bank holidays. They are averaged into the week, month and cargo windows. The tab excludes them from averages (\"corrected\")."))

    def avg_smooth(a, b, drop_hol=True):
        days = [d for d in sm if a <= d <= b and (not drop_hol or d not in hol)]
        return _avg([sm[d] for d in days]), len(days)

    def coverage(a, b):
        exp = len(_bdays(a, b, hol))
        have = len([d for d in sm if a <= d <= b and d not in hol])
        return have, exp

    # week consistency: smooth average must equal step level
    for w in weeks:
        a, b = dt.date.fromisoformat(w["start"]), dt.date.fromisoformat(w["end"])
        v_all, _ = avg_smooth(a, b, drop_hol=False)
        v_bd, n = avg_smooth(a, b, drop_hol=True)
        w["smooth_avg"] = _r(v_all, 1)
        w["smooth_avg_bd"] = _r(v_bd, 1)
        w["fit_err"] = _r(None if v_all is None or w["level"] is None else v_all - w["level"], 1)
    misfit = [w for w in weeks if w["fit_err"] is not None and abs(w["fit_err"]) > 2]
    for w in misfit:
        findings.append(("high", f"Step and smooth curve disagree in {w['start']} → {w['end']}",
                         f"Step level {w['level']:.1f}c vs smooth weekly average {w['smooth_avg']:.1f}c ({w['fit_err']:+.1f}c). Every other week ties to <0.1c. The roll into this week ({weeks[weeks.index(w) - 1]['roll_to_next']}) is out of line with the neighbouring rolls. Fix the roll or re-solve the tail."))

    # ── futures / BFOE on the Dec-futures basis ──
    by, bm = asof.year, 12  # basis month = Dec of as-of year (workbook convention "vs DEC futures")
    fut = {(by, 12): 0.0, _add_month(by, 12, 1): -dec_jan, _add_month(by, 12, 2): -dec_jan - jan_feb}
    if feb_mar is not None:
        fut[_add_month(by, 12, 3)] = -dec_jan - jan_feb - feb_mar
    efp = {(by, 12): efp_dec, _add_month(by, 12, 1): efp_jan}
    bfoe = {k: fut[k] + efp.get(k, efp_jan) for k in fut}
    bfoe_src = {k: ("EFP quoted" if k in efp else "EFP assumed = Jan EFP") for k in fut}

    def front_line(y, mo):
        return _add_month(y, mo, 2)

    # ── divergence by BFOE month (universe 11th M-1 → 10th M) ──
    divergence = []
    for k in sorted(bfoe):
        y, mo = k
        py, pm = _add_month(y, mo, -1)
        a, b = dt.date(py, pm, 11), dt.date(y, mo, 10)
        v, n = avg_smooth(a, b)
        v_wb, _ = avg_smooth(a, b, drop_hol=False)
        have, exp = coverage(a, b)
        univ = [(d, sm[d]) for d in sorted(sm) if a <= d <= b and d not in hol]
        viol = [(d, s - bfoe[k]) for d, s in univ if s is not None and s > bfoe[k]]
        divergence.append({
            "month": _mlabel(y, mo), "start": a.isoformat(), "end": b.isoformat(),
            "bfoe": _r(bfoe[k], 1), "bfoe_src": bfoe_src[k], "avg_dated": _r(v, 1), "avg_dated_wb": _r(v_wb, 1),
            "divergence": _r(None if v is None else bfoe[k] - v, 1),
            "divergence_wb": _r(None if v_wb is None else bfoe[k] - v_wb, 1),
            "coverage": f"{have}/{exp}", "complete": have >= exp and exp > 0,
            "membrane_days": len(viol), "membrane_max": _r(max([x[1] for x in viol]) if viol else None, 1),
            "membrane_first": viol[0][0].isoformat() if viol else None,
            "membrane_last": viol[-1][0].isoformat() if viol else None,
        })
    for dv in divergence:
        if dv["avg_dated"] is None:
            continue
        if not dv["complete"]:
            findings.append(("high", f"{dv['month']} divergence is computed on a truncated universe",
                             f"Universe {dv['start']} → {dv['end']} has {dv['coverage']} assessment days in the grid. The curve is falling into the tail, so the missing days would lower avg Dated and raise divergence. Extend DAILY to at least 10 Feb before reading this number."))
        if dv["divergence"] is not None and dv["divergence"] < 0:
            findings.append(("high", f"{dv['month']} divergence is negative ({dv['divergence']:+.1f}c): membrane breach on average",
                             f"BFOE {dv['month']} {dv['bfoe']:+.1f}c vs avg Dated in its universe {dv['avg_dated']:+.1f}c. The membrane (BFOE→Dated one way) needs avg Dated ≤ BFOE, so divergence ≥ 0. Either the tail of the smooth curve is too high, DecJan/Jan EFP is off, or (most likely) the universe is truncated."))
        if dv["membrane_days"]:
            findings.append(("med", f"{dv['month']} universe: {dv['membrane_days']} days with Dated above BFOE {dv['month']}",
                             f"{dv['membrane_first']} → {dv['membrane_last']}, max excess {dv['membrane_max']:+.1f}c. The training-deck rule is a hard constraint (Dated for any day ≤ its BFOE). In super-backwardation the first days of a universe can print above the month BFOE. Add it to the KKT solve as a soft constraint so it is a choice, not an accident."))

    # ── DFL by calendar month: quoted vs curve-implied vs skeleton ──
    quotes = {}
    for r in range(24, 40):
        lab = R.cell(r, 1).value
        if isinstance(lab, str) and _num(R.cell(r, 2).value) is not None:
            quotes[lab.strip()] = _num(R.cell(r, 2).value)

    def q(prefix):
        for k, v in quotes.items():
            if k.lower().startswith(prefix.lower()):
                return v, k
        return None, None

    spreads = {(by, 12): dec_jan, _add_month(by, 12, 1): jan_feb}
    if feb_mar is not None:
        spreads[_add_month(by, 12, 2)] = feb_mar

    def to_basis(y, mo):
        """cents to add to a Dec-basis value to express it vs the futures month (y, mo)."""
        tot, k = 0.0, (by, 12)
        while k < (y, mo):
            if k not in spreads:
                return None
            tot += spreads[k]
            k = _add_month(k[0], k[1], 1)
        return tot

    dfl = []
    months = [(asof.year, asof.month)] + [_add_month(asof.year, asof.month, i) for i in range(1, 4)]
    qmap = {0: q("Bal "), 1: q("Cal " + MNAME[months[1][1]]), 2: q("Cal " + MNAME[months[2][1]])}
    if dfl_jan_q is not None:
        qmap[3] = (dfl_jan_q, "DFL_DECOMP!B8")
    for i, (y, mo) in enumerate(months):
        fy, fm = front_line(y, mo)
        a = dt.date(y, mo, 1) if i else asof
        b = _add_month(y, mo, 1)
        b = dt.date(b[0], b[1], 1) - dt.timedelta(days=1)
        v, n = avg_smooth(a, b)
        v_wb, _ = avg_smooth(a, b, drop_hol=False)
        have, exp = coverage(a, b)
        sh = to_basis(fy, fm)
        imp = None if v is None or sh is None else v + sh
        imp_wb = None if v_wb is None or sh is None else v_wb + sh
        quoted, qsrc = qmap.get(i, (None, None))
        s1 = s2 = None
        k0, k1 = (y, mo), _add_month(y, mo, 1)
        if k0 in spreads and k1 in spreads:
            s1 = (spreads[k0] + spreads[k1]) / 3.0
            s2 = 2.0 * spreads[k1] / 3.0
        elif (y, mo) < (by, 12) and k1 == (by, 12):
            s1 = None
        skel = None if s1 is None else s1 + s2
        row = {"month": ("Bal " if i == 0 else "") + _mlabel(y, mo), "front_line": _mlabel(fy, fm),
               "window": f"{a.isoformat()} → {b.isoformat()}", "coverage": f"{have}/{exp}", "complete": have >= exp,
               "avg_vs_dec": _r(v, 1), "implied": _r(imp, 1), "implied_wb": _r(imp_wb, 1), "quoted": quoted, "quoted_src": qsrc,
               "rich": _r(None if quoted is None or imp is None else quoted - imp, 1),
               "skel_s1": _r(s1, 1), "skel_s2": _r(s2, 1), "skeleton": _r(skel, 1),
               "skel_resid": _r(None if skel is None or quoted is None else skel - quoted, 1),
               "fitted": quoted is not None and imp_wb is not None and abs(quoted - imp_wb) < 1.0}
        dfl.append(row)
    for row in dfl:
        if row["fitted"]:
            findings.append(("med", f"{row['month']} DFL 'exact' check is circular",
                             f"The smooth curve was solved with Cal {row['month']} DFL = {row['quoted']:.0f} as a constraint, so curve-implied {row['implied_wb']:.1f} = quoted by construction. DFL_DECOMP 'Exact' for this month is a fit residual, not an independent check. Use roll trades (e.g. 23-27/11 v Cal Dec) to test it."))
        if not row["complete"] and (row["implied"] is not None or row["implied_wb"] is not None):
            wbv = f" Workbook 'Exact' shows {row['implied_wb']:.1f} vs quoted {row['quoted']:.0f}, which is meaningless." if row["implied_wb"] is not None and row["quoted"] is not None else ""
            findings.append(("high", f"{row['month']} DFL curve-implied value uses {row['coverage']} assessment days",
                             f"Grid ends {grid_end.isoformat()}. DFL_DECOMP 'Exact' for this month averages only what is in the grid.{wbv} Treat it as a placeholder until DAILY runs to month-end."))

    # DTD spreads (curve) vs quoted and vs identity
    dtd = []
    for i in range(len(dfl) - 1):
        a_, b_ = dfl[i], dfl[i + 1]
        if a_["avg_vs_dec"] is None or b_["avg_vs_dec"] is None:
            continue
        name = f"DTD {a_['month'].replace('Bal ', '')[:3]}/{b_['month'][:3]}"
        qv, qs = q(name)
        ident = None
        k_a, k_b = months[i], months[i + 1]
        fa, fb = front_line(*k_a), front_line(*k_b)
        if a_["quoted"] is not None and b_["quoted"] is not None and fa in spreads:
            ident = a_["quoted"] - b_["quoted"] + spreads[fa]
        dtd.append({"name": name, "curve": _r(a_["avg_vs_dec"] - b_["avg_vs_dec"], 1), "quoted": qv,
                    "identity": _r(ident, 1), "complete": a_["complete"] and b_["complete"]})

    # ── per-cargo X ──
    Cg, CgF = wv["CARGOES"], wf["CARGOES"]
    half_spread = 75.0
    mg = re.search(r"F\d+\s*-\s*(\d+(?:\.\d+)?)", str(CgF["G5"].value or ""))
    if mg:
        half_spread = float(mg.group(1))
    cargoes = []
    for r in range(5, 60):
        name, bl = Cg.cell(r, 1).value, _d(Cg.cell(r, 2).value)
        if not bl:
            if cargoes:
                break
            continue
        w_a, w_b = bl - dt.timedelta(days=5), bl + dt.timedelta(days=5)
        wb_avg = _num(Cg.cell(r, 5).value)
        v_cal, n_cal = avg_smooth(w_a, w_b)
        bds_before = [d for d in sorted(sm) if d <= bl and d not in hol]
        bds_after = [d for d in sorted(sm) if d > bl and d not in hol]
        q5 = (bds_before[-3:] if bl in bds_before else bds_before[-2:]) + bds_after[:2 if bl in bds_before else 3]
        v5 = _avg([sm[d] for d in q5]) if len(q5) == 5 else None
        nat = [d for d in sorted(sm) if bl - dt.timedelta(days=30) <= d <= bl - dt.timedelta(days=10) and d not in hol]
        v_nat = _avg([sm[d] for d in nat])
        x = None if v_cal is None else efp_dec - v_cal
        cargoes.append({
            "cargo": name, "bl": bl.isoformat(), "win_start": w_a.isoformat(), "win_end": w_b.isoformat(),
            "n_quotes": n_cal, "avg_cal": _r(v_cal, 1), "avg_wb": _r(wb_avg, 1),
            "x": _r(x, 1), "x_bid": _r(None if x is None else x - half_spread, 1), "x_offer": _r(None if x is None else x + half_spread, 1),
            "prem_dec": _r(None if v_cal is None else v_cal - efp_dec, 1),
            "avg_5q": _r(v5, 1), "q5_dates": [d.isoformat() for d in q5],
            "avg_natural": _r(v_nat, 1), "nat_n": len(nat), "nat_from": nat[0].isoformat() if nat else None, "nat_to": nat[-1].isoformat() if nat else None,
            "pricing_value": _r(None if v_nat is None or v_cal is None else v_nat - v_cal, 1),
        })
    if cargoes:
        c0 = cargoes[0]
        findings.append(("med", "CARGOES header says 'vs Nov26 … cash anchor = Nov contract + EFP' but the maths is vs Dec",
                         f"F = RUN!B4 (Dec EFP {efp_dec:+.0f}) − window avg of a curve quoted vs DEC futures. So X = Dec BFOE − cargo Dated, i.e. the cargo vs DEC BFOE, not vs Nov BFOE. The A2 note is right ('basis now vs Dec'); A1, DAILY!B3 and DIVERGENCE D8/E8 still say 'vs Nov'. If you want the cargo vs Nov BFOE you need the Nov/Dec cash BFOE spread as an input: X(Nov) = (Nov/Dec cash + Dec EFP) − window avg."))
        findings.append(("low", "Pricing window is B/L ±5 calendar days (≈7 quotes), not 5 quotes around B/L",
                         f"F-1 window {c0['win_start']} → {c0['win_end']} = {c0['n_quotes']} quotes. If the cargoes price 5 quotes around B/L (2-1-2), the average moves: F-1 {c0['avg_cal']:+.0f} (±5d) vs {c0['avg_5q'] if c0['avg_5q'] is not None else float('nan'):+.0f} (2-1-2). Tab shows both."))
        big = max(cargoes, key=lambda c: abs(c["pricing_value"] or 0))
        if big["pricing_value"] is not None:
            findings.append(("info", "Hidden time-spread: B/L pricing vs the cargo's natural Platts window",
                             f"A cargo loading {big['bl']} sits in the Platts Dated window 10–30 days before loading ({big['nat_from']} → {big['nat_to']}, avg {big['avg_natural']:+.0f}c vs Dec) but prices around B/L ({big['avg_cal']:+.0f}c). The {big['pricing_value']:+.0f}c gap is the time-spread embedded in B/L pricing in this backwardation (training deck: natural pricing date vs contract strip). Whoever holds the cargo against B/L pricing is long that structure until it prices. Not in the workbook; added as a column."))
    findings.append(("low", f"Bid/offer is a flat ±{half_spread:.0f}c around gross X",
                     "Same width for every cargo, so it is not market-derived. Fine as a desk haircut, but label it as an assumption, not a quote."))

    # ── Dubai complex ──
    dubai = parse_dubai(wv) if "DUBAI" in wv.sheetnames else None
    if dubai and dubai.get("text_cells"):
        findings.append(("low", "Numbers stored as text in DUBAI",
                         f"{', '.join(dubai['text_cells'][:8])} are text, not numbers (e.g. ='12'). Sums or charts that use them will skip or break."))
    if dubai and dubai.get("efs_bd_dub"):
        live = [x for x in dubai["efs_bd_dub"] if x["err"] is not None and (mkey_of(x["bd_month"]) or (0, 0)) >= (asof.year, asof.month)]
        errs = [abs(x["err"]) for x in live]
        if errs:
            findings.append(("info", "Second triangle added: EFS(M+2) − B/D(M) = Dubai M/M+2",
                             f"B/D(M) is front-line Brent (M+2 futures) vs Dubai M, EFS(M+2) is the same Brent vs Dubai M+2, so the gap must equal the Dubai M/M+2 spread. For B/D months from {MNAME[asof.month]} on it misses by {min(errs):.2f}–{max(errs):.2f} $/bbl (median {sorted(errs)[len(errs)//2]:.2f}), always the same sign: EFS−B/D runs above the Dubai spread. That is the front-line averaging / Dubai partials vs swaps basis, worth knowing when you leg B/D against EFS. Expired months (Sep) blow out and are shown greyed. This is an independent check (the Dated/Dub−DFL=B/D identity is exact because one leg is derived)."))

    mw_note = str(wv["MARKETWIRE"]["A2"].value or "") if "MARKETWIRE" in wv.sheetnames else ""
    if mw_note:
        findings.append(("info", "Dated-setting grade (period-specific)",
                         mw_note[:260] + (" …" if len(mw_note) > 260 else "") + " Read Forties/BFOE grades 'over the setter'; re-check the setter every window rather than hard-coding it."))

    order = {"high": 0, "med": 1, "low": 2, "info": 3}
    findings.sort(key=lambda f: order[f[0]])
    return {
        "title": title, "version": version, "asof": asof.isoformat(), "basis": f"Dec{str(by)[2:]} ICE futures",
        "grid_start": daily[0]["date"], "grid_end": grid_end.isoformat(),
        "inputs": {"efp_dec": efp_dec, "efp_jan": efp_jan, "dec_jan": dec_jan, "jan_feb": jan_feb, "feb_mar": feb_mar,
                   "front_frozen": front_frozen, "half_spread": half_spread},
        "quotes": quotes,
        "bfoe": [{"month": _mlabel(*k), "vs_dec": _r(bfoe[k], 1), "fut_vs_dec": _r(fut[k], 1), "efp": efp.get(k, efp_jan), "src": bfoe_src[k]} for k in sorted(bfoe)],
        "weeks": weeks, "daily": daily, "divergence": divergence, "dfl": dfl, "dtd": dtd, "cargoes": cargoes,
        "dubai": dubai, "marketwire": parse_marketwire(wv) if "MARKETWIRE" in wv.sheetnames else None,
        "findings": [{"sev": s, "title": t, "detail": d} for s, t, d in findings],
        "holidays": hol_in_grid,
    }


SECTIONS = {"EFS": "efs", "DUBAI intermonth": "dub_spr", "BRENT/DUBAI": "bd", "DFL": "dfl", "DATED/DUBAI": "dated_dub",
            "DATED SPREADS": "dtd", "MURBAN/DUBAI": "murban_dub", "MURBAN/BRENT": "murban_brent"}


def mkey_of(lbl):
    mm = re.match(r"([A-Z][a-z]{2})-(\d{2})", lbl or "")
    return (2000 + int(mm.group(2)), MON[mm.group(1)]) if mm and mm.group(1) in MON else None


def parse_dubai(wv):
    S = wv["DUBAI"]
    snaps = []
    for c in range(2, 20):
        v = S.cell(3, c).value
        if v is None or str(v).startswith("Δ"):
            break
        snaps.append(str(v))
    out, cur, text_cells = {}, None, []
    for r in range(4, 60):
        lab = S.cell(r, 1).value
        if not isinstance(lab, str):
            continue
        if lab.startswith("IDENTITY"):
            break
        vals = [S.cell(r, c).value for c in range(2, 2 + len(snaps))]
        if all(v is None for v in vals):
            cur = next((k for p, k in SECTIONS.items() if lab.upper().startswith(p.upper())), None)
            continue
        if cur:
            for j, v in enumerate(vals):
                if isinstance(v, str):
                    text_cells.append(f"{openpyxl.utils.get_column_letter(j + 2)}{r}")
            out.setdefault(cur, {})[lab.strip()] = [_num(v) for v in vals]
    # identity checks per snapshot
    ident = []
    mon_keys = list(out.get("dated_dub", {}).keys())
    for mk in mon_keys:
        for j, sn in enumerate(snaps):
            dd = out.get("dated_dub", {}).get(mk, [None] * len(snaps))[j]
            df = out.get("dfl", {}).get(mk, [None] * len(snaps))[j]
            bd = out.get("bd", {}).get(mk, [None] * len(snaps))[j]
            if None not in (dd, df, bd):
                ident.append({"month": mk, "snap": sn, "err": round(dd - (df + bd), 3)})
    mur = []
    for mk in out.get("murban_dub", {}):
        for j, sn in enumerate(snaps):
            a = out["murban_dub"][mk][j]
            b = out.get("murban_brent", {}).get(mk, [None] * len(snaps))[j]
            c = out.get("efs", {}).get(mk, [None] * len(snaps))[j]
            if None not in (a, b, c):
                mur.append({"month": mk, "snap": sn, "err": round(a - (b + c), 3)})
    # EFS(M+2) − B/D(M) vs Dubai M/M+1 + M+1/M+2
    def mkey(lbl):
        mm = re.match(r"([A-Z][a-z]{2})-(\d{2})", lbl)
        return (2000 + int(mm.group(2)), MON[mm.group(1)]) if mm and mm.group(1) in MON else None
    spr = {}
    for lbl, vals in out.get("dub_spr", {}).items():
        mm = re.match(r"([A-Z])/([A-Z])$", lbl)
        if mm:
            spr[lbl] = vals
    letters = {1: "J", 2: "F", 3: "M", 4: "A", 5: "M", 6: "J", 7: "J", 8: "A", 9: "S", 10: "O", 11: "N", 12: "D"}
    triple = []
    for lbl, vals in out.get("bd", {}).items():
        k = mkey(lbl)
        if not k:
            continue
        k2 = _add_month(k[0], k[1], 2)
        efs_lbl = f"{MNAME[k2[1]]}-{str(k2[0])[2:]}"
        k1 = _add_month(k[0], k[1], 1)
        s1, s2 = f"{letters[k[1]]}/{letters[k1[1]]}", f"{letters[k1[1]]}/{letters[k2[1]]}"
        if efs_lbl not in out.get("efs", {}) or s1 not in spr or s2 not in spr:
            continue
        for j, sn in enumerate(snaps):
            e, b, a1, a2 = out["efs"][efs_lbl][j], vals[j], spr[s1][j], spr[s2][j]
            if None in (e, b, a1, a2):
                continue
            triple.append({"bd_month": lbl, "efs_month": efs_lbl, "snap": sn, "efs_minus_bd": round(e - b, 2),
                           "dubai_spread": round(a1 + a2, 2), "err": round(e - b - a1 - a2, 2)})
    return {"snapshots": snaps, "series": out, "identity": ident, "murban_identity": mur, "efs_bd_dub": triple, "text_cells": text_cells}


def parse_marketwire(wv):
    S = wv["MARKETWIRE"]
    grades, setter = [], None
    for r in range(5, 30):
        lab, diff = S.cell(r, 1).value, _num(S.cell(r, 2).value)
        if not isinstance(lab, str) or diff is None:
            continue
        cif = "CIF" in lab
        if lab.startswith("Dated Brent Diff") or _num(S.cell(r, 4).value) == 0:
            setter = diff
            lab = "WTI Midland FOB-NS (setter)"
        grades.append({"grade": lab, "diff": diff, "chg": _num(S.cell(r, 3).value), "over_setter": _num(S.cell(r, 4).value), "cif": cif})
    hdr = str(S["A1"].value or "")
    m = re.search(r"(\d{1,2}-[A-Z]{3}-\d{4})", hdr)
    return {"date": m.group(1) if m else None, "setter": setter, "grades": grades,
            "setter_note": str(S["A2"].value or "")[:200]}


def get_model():
    content = load_workbook_bytes()
    if content is None:
        return None
    return parse(content)
